"""The spreadsheet toggle and the player page both read season aggregates."""

from datetime import datetime, timezone

from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.config import Settings
from app.db.base import Base
from app.db.models import Player, PlayerSeasonAggregate
from app.db.session import get_db
from app.main import app
from app.services.refresh import refresh_data

from fakes import SquadClient


def _aggregate(season, points, **overrides):
    values = dict(
        player_code=5001, season=season, fixtures=38, appearances=34, minutes=2900,
        starts=33, starts_derived=False, points=points, goals=10, assists=8,
        clean_sheets=9, bonus=20, appearance_points=points,
        appearance_points_sq=points * 8, blanks=12, hauls=5, price_min=8.0,
        price_max=8.6, position="MID", team_name="Test FC",
        computed_at=datetime(2026, 9, 1, tzinfo=timezone.utc),
    )
    values.update(overrides)
    return PlayerSeasonAggregate(**values)


def _client(tmp_path, *, with_history=True):
    url = f"sqlite:///{tmp_path / 'season-history.db'}"
    engine = create_engine(url, connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    Session = sessionmaker(engine, expire_on_commit=False)
    with Session() as db:
        refresh_data(db, Settings(database_url=url, current_season="2026/27"), SquadClient())
        # Player 10 has a past record; player 11 is new to the league.
        db.get(Player, 10).code = 5001
        db.get(Player, 11).code = 5002
        if with_history:
            db.add_all([
                _aggregate("2025/26", 211),
                _aggregate("2024/25", 203, starts_derived=True),
                # A current-season aggregate must never be shown as history.
                _aggregate("2026/27", 777),
            ])
        db.commit()

    def override_db():
        with Session() as db:
            yield db

    app.dependency_overrides[get_db] = override_db
    return TestClient(app)


def test_spreadsheet_hides_history_until_toggled(tmp_path):
    try:
        html = _client(tmp_path).get("/spreadsheet?view=forward").text
    finally:
        app.dependency_overrides.clear()

    assert "Consistency" not in html
    assert 'href="/spreadsheet?view=forward&amp;history=1"' in html
    assert "Show past seasons" in html


def test_spreadsheet_history_adds_derived_and_per_season_columns(tmp_path):
    try:
        html = _client(tmp_path).get("/spreadsheet?history=1").text
    finally:
        app.dependency_overrides.clear()

    for heading in ("Seasons", "Avg P/90", "Spread", "Consistency", "GW SD",
                    "Blank %", "Haul %", "Hist Start %", "Durability",
                    "25/26", "24/25"):
        assert f'<span class="column-label">{heading}</span>' in html
    assert ">26/27<" not in html and "777" not in html
    assert 'data-sort-value="211"' in html and 'data-sort-value="203"' in html
    assert "Steady" in html and "High" in html
    # The toggle survives switching tab and applying filters.
    assert 'href="/spreadsheet?view=rotation&amp;history=1"' in html
    assert '<input type="hidden" name="history" value="1">' in html
    assert "Hide past seasons" in html


def test_spreadsheet_history_explains_a_missing_build(tmp_path):
    try:
        html = _client(tmp_path, with_history=False).get("/spreadsheet?history=1").text
    finally:
        app.dependency_overrides.clear()

    assert "build-season-aggregates" in html


def test_player_page_shows_season_history_or_says_there_is_none(tmp_path):
    try:
        client = _client(tmp_path)
        established = client.get("/players/10").text
        newcomer = client.get("/players/11").text
    finally:
        app.dependency_overrides.clear()

    assert "Season history" in established
    assert "2025/26" in established and "2024/25" in established
    assert "211" in established and "777" not in established
    assert "inferred from minutes" in established
    assert "No past-season record" in newcomer


def test_exports_carry_perfect_pick_and_past_seasons(tmp_path):
    import csv
    import io

    from openpyxl import load_workbook

    try:
        client = _client(tmp_path)
        text = client.get("/exports/current.csv").content.decode("utf-8-sig")
        workbook = load_workbook(
            io.BytesIO(client.get("/exports/current.xlsx").content), read_only=True
        )
    finally:
        app.dependency_overrides.clear()

    records = list(csv.DictReader(io.StringIO(text)))
    header = list(records[0])
    assert header.index("Value") < header.index("Perfect Pick") < header.index("Reliable Value")
    established = next(row for row in records if row["2025/26 Points"] == "211")
    assert established["Past Seasons"] == "2" and established["Consistency"] == "Steady"
    assert established["2024/25 Points"] == "203"
    assert "2026/27 Points" not in header
    # No record exports as empty cells, never zeros.
    newcomer = next(row for row in records if row["Past Seasons"] == "")
    assert newcomer["2025/26 Points"] == "" and newcomer["Durability"] == ""
    assert float(established["Perfect Pick"]) > 0

    assert "Past Seasons" in workbook.sheetnames
    past_header = [cell.value for cell in next(workbook["Past Seasons"].iter_rows())]
    assert past_header[:4] == ["Player", "Team", "Position", "Price"]
    assert "Consistency" in past_header and "2025/26 Points" in past_header


def test_player_page_leads_with_perfect_pick(tmp_path):
    try:
        html = _client(tmp_path).get("/players/10").text
    finally:
        app.dependency_overrides.clear()

    assert html.index("<span>Perfect Pick</span>") < html.index("<span>Raw value</span>")
    assert "Ranked " in html
