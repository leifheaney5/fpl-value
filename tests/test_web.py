from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from fastapi.testclient import TestClient
from io import BytesIO
from openpyxl import load_workbook
from types import SimpleNamespace
from pathlib import Path
import re

from app.db.base import Base
from app.db.session import get_db
from app.main import app
from app.config import Settings
from app.web.auth import safe_next_path, valid_credentials
from app.web import routes


def responsive_rules_at_viewport(css, *, viewport_width):
    """Return the first max-width media block that applies at this width."""
    for match in re.finditer(r"@media\s*\(max-width:\s*(\d+)px\)\s*\{", css):
        if viewport_width > int(match.group(1)):
            continue
        depth = 1
        for index in range(match.end(), len(css)):
            if css[index] == "{":
                depth += 1
            elif css[index] == "}":
                depth -= 1
                if depth == 0:
                    return css[match.end():index]
    raise AssertionError(f"No responsive CSS block applies at {viewport_width}px")


def test_web_routes_health_exports_and_new_pages(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'web.db'}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    Session = sessionmaker(engine, expire_on_commit=False)

    def override_db():
        with Session() as db:
            yield db

    app.dependency_overrides[get_db] = override_db
    try:
        client = TestClient(app)
        assert client.get("/health").status_code == 200
        for path in ["/", "/players", "/spreadsheet", "/recommendation", "/transfers", "/movers", "/compare", "/diagnostics", "/schema", "/settings", "/differentials", "/transfer-market", "/templates", "/fixtures", "/performance"]:
            assert client.get(path).status_code == 200
        assert "No upcoming fixtures are available yet." in client.get("/fixtures").text
        assert "No scored fixtures are available yet." in client.get("/performance").text
        # Personal routes stay closed to anonymous visitors even in demo mode.
        assert client.get("/my-team", follow_redirects=False).status_code == 303
        assert client.get("/players?position=MID").url.path == "/spreadsheet"
        assert client.get("/players?max_price=&min_minutes=&max_ownership=not-a-number").status_code == 200
        login_page = client.get("/login")
        assert login_page.status_code == 200
        assert client.post("/login", data={"username": "x", "password": "y", "csrf_token": "bad"}).status_code == 401
        csv_response = client.get("/exports/current.csv")
        assert csv_response.status_code == 200
        assert b"Player" in csv_response.content
        xlsx_response = client.get("/exports/current.xlsx")
        assert xlsx_response.headers["content-type"].startswith("application/vnd.openxmlformats")
        workbook = load_workbook(BytesIO(xlsx_response.content), read_only=True)
        assert {"Dashboard", "Value Rankings", "Forward Value", "Rotation Risk", "Movers", "Fixtures", "Schema Changes", "Guide"}.issubset(workbook.sheetnames)
    finally:
        app.dependency_overrides.clear()


def test_header_regions_center_navigation_and_preserve_responsive_collapse(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'header.db'}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    Session = sessionmaker(engine, expire_on_commit=False)

    def override_db():
        with Session() as db:
            yield db

    app.dependency_overrides[get_db] = override_db
    try:
        html = TestClient(app).get("/").text
        assert 'class="site-header__brand ' in html
        assert 'class="site-header__nav"' in html
        assert 'class="site-header__actions"' in html
        assert html.index("site-header__brand") < html.index("site-header__nav") < html.index("site-header__actions")
        assert 'href="/fixtures"' in html
        assert 'href="/performance"' in html
        assert 'href="/settings"' in html

        css = (Path(routes.__file__).parent.parent / "static" / "app.css").read_text()
        assert ".site-header {" in css
        assert "grid-template-columns: minmax(0, 1fr) auto minmax(0, 1fr)" in css
        assert ".site-header__nav {" in css
        assert "@media (max-width: 1200px)" in css
        assert "grid-template-columns: minmax(0, 1fr) auto" in responsive_rules_at_viewport(css, viewport_width=1024)
        narrow_rules = responsive_rules_at_viewport(css, viewport_width=600)
        assert re.search(r"\.site-header__nav\s*\{\s*display:\s*none;\s*\}", narrow_rules)
        assert "grid-template-columns: minmax(0, 1fr) auto" in narrow_rules
    finally:
        app.dependency_overrides.clear()


def test_credentials_use_constant_time_path_and_safe_redirect():
    settings = Settings(app_username="admin", app_password="secret")
    assert settings.credentials_configured
    assert valid_credentials(settings, "admin", "secret")
    assert not valid_credentials(settings, "admin", "wrong")
    # Fail closed: with no credentials configured nothing authenticates.
    assert not valid_credentials(Settings(), "admin", "secret")
    assert safe_next_path("/players") == "/players"
    assert safe_next_path("https://example.invalid") == "/"
    assert safe_next_path("//example.invalid") == "/"


def test_comparison_scale_ranks_values_and_keeps_missing_values_neutral():
    assert routes.comparison_scale([10, 20, 30]) == [
        "compare-scale-0",
        "compare-scale-2",
        "compare-scale-4",
    ]
    assert routes.comparison_scale([10, 20, 30], lower_is_better=True) == [
        "compare-scale-4",
        "compare-scale-2",
        "compare-scale-0",
    ]
    assert routes.comparison_scale([10, None, 30]) == [
        "compare-scale-0",
        "compare-scale-neutral",
        "compare-scale-4",
    ]
    assert routes.comparison_scale([10, 10]) == [
        "compare-scale-neutral",
        "compare-scale-neutral",
    ]


def test_compare_context_uses_metric_aware_scales(monkeypatch):
    def row(player_id, *, price, value, risk):
        snapshot = SimpleNamespace(
            price=price,
            total_points=10 if player_id == 1 else 20,
            value=value,
            reliable_value=value,
            forward_value=value,
            projected_points_5=value,
            rotation_risk=risk,
            points_per_game=value,
            points_per_90=value,
            points_per_start=value,
            minutes=900 if player_id == 1 else 1_000,
            start_rate=value,
            expected_minutes=value,
            ownership=value,
            average_fixture_difficulty=risk,
            metric_status={},
        )
        return {"player": SimpleNamespace(id=player_id), "snapshot": snapshot}

    rows = [
        row(1, price=6.0, value=1.0, risk=4.0),
        row(2, price=8.0, value=2.0, risk=2.0),
    ]
    monkeypatch.setattr(routes, "filtered_players", lambda db, season, sort: rows)
    monkeypatch.setattr(routes.templates, "TemplateResponse", lambda **kwargs: kwargs["context"])

    context = routes.compare(
        SimpleNamespace(), db=object(), settings=Settings(current_season="2026/27"), ids=[1, 2]
    )
    scales = {metric["label"]: metric["classes"] for metric in context["comparison_metrics"]}

    assert scales["Price"] == ["compare-scale-4", "compare-scale-0"]
    assert scales["Raw value"] == ["compare-scale-0", "compare-scale-4"]
    assert scales["Rotation risk"] == ["compare-scale-0", "compare-scale-4"]


def test_page_asset_versions_match_served_contents():
    from hashlib import sha256

    client = TestClient(app)
    request = SimpleNamespace(state=SimpleNamespace(
        can_access_personal=False, authenticated=False, sign_in_available=False,
    ))
    html = routes.templates.env.get_template("base.html").render(request=request)
    assets = re.findall(r'(?:src|href)="(/static/app\.(?:css|js)\?v=([a-f0-9]{12}))"', html)
    assert len(assets) == 2
    for url, version in assets:
        response = client.get(url)
        assert response.status_code == 200
        assert sha256(response.content).hexdigest()[:12] == version
    assert f'/static/app.css?v={routes.templates.env.globals["asset_versions"]["app.css"]}' in client.get("/login").text


def test_column_help_glossary_covers_every_table_header_key():
    from jinja2 import nodes
    from app.web.column_help import COLUMN_HELP

    template_dir = Path("app/templates")
    header_keys = set()
    for template_path in template_dir.glob("*.html"):
        source = template_path.read_text(encoding="utf-8")
        header_keys.update(re.findall(r'info_header\([^,]+,\s*"([^"]+)"', source))
        # Sortable tables supply their help keys as (label, key) tuples.
        tree = routes.templates.env.parse(source)
        for loop in tree.find_all(nodes.For):
            if isinstance(loop.target, nodes.Tuple) and [item.name for item in loop.target.items] == ["label", "key"]:
                header_keys.update(item.items[1].value for item in loop.iter.items)

    header_keys.update(name for _, name, _ in routes._COMPARE_METRICS)

    assert len(header_keys) >= 30
    assert header_keys <= COLUMN_HELP.keys()
    assert all(COLUMN_HELP[key].strip() for key in header_keys)


def test_every_table_bearing_template_uses_shared_column_help():
    template_dir = Path("app/templates")

    table_templates = []
    for template_path in template_dir.glob("*.html"):
        source = template_path.read_text(encoding="utf-8")
        if "<th" in source:
            table_templates.append(template_path.name)
            assert "info_header" in source, template_path.name
            for header in re.findall(r"<th\b[^>]*>(.*?)</th>", source, re.DOTALL):
                assert "info_header(" in header, (template_path.name, header)

    assert len(table_templates) >= 15


def test_column_help_has_hover_and_keyboard_focus_states():
    css = (Path(routes.__file__).parent.parent / "static" / "app.css").read_text(encoding="utf-8")

    assert ".column-heading:hover .column-tooltip" in css
    assert '.column-heading[tabindex="0"]:focus-visible .column-tooltip' in css
    assert ".column-heading[tabindex=\"0\"]" in css


def test_info_header_renders_explanation_and_keyboard_target():
    template = routes.templates.env.from_string(
        '{% from "_macros.html" import info_header %}{{ info_header("Price", "price") }}'
    )

    rendered = template.render()

    assert 'class="column-heading column-heading--focusable"' in rendered
    assert 'tabindex="0"' in rendered
    assert "current FPL price, shown in millions of pounds" in rendered
    assert 'role="tooltip"' in rendered


def test_dashboard_passes_its_loaded_rows_to_my_team(monkeypatch):
    rows = [{"player": object(), "snapshot": object()}]
    captured_rows = []
    monkeypatch.setattr(routes, "dashboard_data", lambda db, season: {"rows": rows})
    monkeypatch.setattr(
        routes,
        "data_status",
        lambda db, season: {"state": "never_refreshed", "season": season},
    )
    monkeypatch.setattr(routes, "linked_team_data", lambda db, client, settings, *, rows=None: captured_rows.append(rows))
    monkeypatch.setattr(routes.templates, "TemplateResponse", lambda **kwargs: kwargs["context"])
    settings = Settings(access_mode="local", fpl_entry_id=123)

    context = routes.dashboard(SimpleNamespace(session={}), db=object(), settings=settings)

    assert captured_rows == [rows]
    assert context["my_team"] is None


def test_recommendation_passes_its_loaded_rows_to_my_team(monkeypatch):
    rows = [{"player": object(), "snapshot": object()}]
    captured_rows = []
    monkeypatch.setattr(routes, "latest_rows", lambda db, season: rows)
    monkeypatch.setattr(routes, "recommend_team_cached", lambda candidate_rows, budget, strategy: None)
    monkeypatch.setattr(routes, "compare_recommendations", lambda candidate_rows, budget: [])
    monkeypatch.setattr(routes, "linked_team_data", lambda db, client, settings, *, rows=None: captured_rows.append(rows))
    monkeypatch.setattr(routes.templates, "TemplateResponse", lambda **kwargs: kwargs["context"])
    settings = Settings(access_mode="local", fpl_entry_id=123)

    context = routes.recommendation_page(SimpleNamespace(session={}), db=object(), settings=settings)

    assert captured_rows == [rows]
    assert context["my_team"] is None
