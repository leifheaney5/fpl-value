from __future__ import annotations

import csv
import io
from datetime import date, datetime, time

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from sqlalchemy import select
from app.db.models import Fixture, SchemaChange

from app.services.queries import latest_rows
from app.services.season_history import career_summaries, stored_seasons


EXPORT_COLUMNS = [
    ("Player", lambda row: row["player"].full_name),
    ("Team", lambda row: row["team"].name),
    ("Position", lambda row: row["player"].position_short),
    ("Price", lambda row: row["snapshot"].price),
    ("Points", lambda row: row["snapshot"].total_points),
    ("Value", lambda row: row["snapshot"].value),
    ("Perfect Pick", lambda row: row["snapshot"].pick_score),
    ("Perfect Pick Rank", lambda row: row["snapshot"].pick_rank),
    ("Reliable Value", lambda row: row["snapshot"].reliable_value),
    ("Forward Value", lambda row: row["snapshot"].forward_value),
    ("Projected Points", lambda row: row["snapshot"].projected_points_5),
    ("Expected Minutes", lambda row: row["snapshot"].expected_minutes),
    ("Rotation Risk", lambda row: row["snapshot"].rotation_risk),
    ("Rotation Tier", lambda row: row["snapshot"].rotation_tier),
    ("PPG", lambda row: row["snapshot"].points_per_game),
    ("Points/90", lambda row: row["snapshot"].points_per_90),
    ("Points/Minute", lambda row: row["snapshot"].points_per_minute),
    ("Start Rate", lambda row: row["snapshot"].start_rate),
    ("Minutes", lambda row: row["snapshot"].minutes),
    ("Ownership", lambda row: row["snapshot"].ownership),
    ("1D Value Change", lambda row: row["history"]["1D"]["delta_value"]),
    ("1D Value Direction", lambda row: row["history"]["1D"]["value_direction"]),
    ("1D Price Change", lambda row: row["history"]["1D"]["delta_price"]),
    ("1D Price Direction", lambda row: row["history"]["1D"]["price_direction"]),
    ("7D Value Change", lambda row: row["history"]["7D"]["delta_value"]),
    ("7D Rank Change", lambda row: row["history"]["7D"]["delta_rank"]),
    ("30D Value Change", lambda row: row["history"]["30D"]["delta_value"]),
]


def _excel_value(value):
    """Convert timezone-aware dates to Excel-compatible naive values."""
    if isinstance(value, (datetime, time)) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    return value


def _career(name: str, digits: int | None = None):
    def getter(row):
        value = (row.get("career") or {}).get(name)
        return round(value, digits) if digits is not None and value is not None else value
    return getter


def _season_points(season: str):
    return lambda row: ((row.get("career") or {}).get("points_by_season") or {}).get(season)


def _history_columns(
    db: Session, season: str, rows: list[dict], sample_minutes: int
) -> list[tuple[str, object]]:
    """Attach each row's past-season summary and return the columns that read it.

    A player with no record, or a measure withheld for a thin sample, exports as
    an empty cell, in keeping with the rule that an empty cell is not a zero.
    """
    careers = career_summaries(
        db,
        [row["player"].code for row in rows],
        exclude_season=season,
        sample_minutes=sample_minutes,
    )
    for row in rows:
        row["career"] = careers.get(row["player"].code)
    return [
        ("Past Seasons", _career("season_count")),
        ("Qualifying Seasons", _career("qualifying_seasons")),
        ("Avg P/90 (Past)", _career("mean_p90", 2)),
        ("P/90 Spread", _career("p90_spread", 2)),
        ("Consistency", _career("consistency")),
        ("GW Points SD", _career("gw_sd", 2)),
        ("Blank %", _career("blank_rate", 1)),
        ("Haul %", _career("haul_rate", 1)),
        ("Past Start %", _career("start_rate", 1)),
        ("Starts Estimated", _career("starts_estimated")),
        ("Past Minutes Share %", _career("minutes_share", 1)),
        ("Durability", _career("durability")),
    ] + [
        (f"{past} Points", _season_points(past))
        for past in stored_seasons(db, exclude_season=season)
    ]


def csv_bytes(db: Session, season: str, sample_minutes: int = 900) -> bytes:
    rows = latest_rows(db, season)
    columns = EXPORT_COLUMNS + _history_columns(db, season, rows, sample_minutes)
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([column for column, _ in columns])
    for row in rows:
        writer.writerow([getter(row) for _, getter in columns])
    return buffer.getvalue().encode("utf-8-sig")


def xlsx_bytes(db: Session, season: str, sample_minutes: int = 900) -> bytes:
    rows = latest_rows(db, season)
    history_columns = _history_columns(db, season, rows, sample_minutes)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dashboard"

    header_fill = PatternFill("solid", fgColor="17365D")
    header_font = Font(color="FFFFFF", bold=True)
    green = PatternFill("solid", fgColor="C6EFCE")
    yellow = PatternFill("solid", fgColor="FFF2CC")
    orange = PatternFill("solid", fgColor="F4B183")
    red = PatternFill("solid", fgColor="F4CCCC")
    gray = PatternFill("solid", fgColor="E7E6E6")

    sheet.append(["Metric", "Value"])
    sheet.append(["Players", len(rows)])
    sheet.append(["Ranked players", sum(1 for row in rows if row["snapshot"].value_rank)])

    def add_table(title: str, columns: list[tuple[str, object]], values: list[object]) -> None:
        target = workbook.create_sheet(title)
        target.append([column for column, _ in columns])
        for cell in target[1]:
            cell.fill = header_fill
            cell.font = header_font
        for row in values:
            target.append([_excel_value(getter(row)) for _, getter in columns])
            tier_index = next((i for i, (name, _) in enumerate(columns, start=1) if name in {"Rotation Tier", "Tier"}), None)
            if tier_index:
                tier_cell = target.cell(target.max_row, tier_index)
                tier_cell.fill = {"Low": green, "Moderate": yellow, "High": orange, "Very High": red}.get(str(tier_cell.value), gray)
            for index, (name, _) in enumerate(columns, start=1):
                if "Change" in name or "Direction" in name:
                    value = target.cell(target.max_row, index).value
                    if isinstance(value, (int, float)):
                        target.cell(target.max_row, index).fill = green if value > 0 else red if value < 0 else gray
        target.freeze_panes = "A2"
        target.auto_filter.ref = target.dimensions
        for index, column in enumerate(columns, start=1):
            target.column_dimensions[get_column_letter(index)].width = max(12, min(28, len(column[0]) + 4))

    add_table("Value Rankings", EXPORT_COLUMNS, rows)
    add_table("Past Seasons", EXPORT_COLUMNS[:4] + history_columns, rows)
    add_table("Forward Value", [("Player", lambda r: r["player"].full_name), ("Position", lambda r: r["player"].position_short), ("Projected Points", lambda r: r["snapshot"].projected_points_5), ("Forward Value", lambda r: r["snapshot"].forward_value), ("Expected Minutes", lambda r: r["snapshot"].expected_minutes)], rows)
    add_table("Rotation Risk", [("Player", lambda r: r["player"].full_name), ("Risk", lambda r: r["snapshot"].rotation_risk), ("Tier", lambda r: r["snapshot"].rotation_tier), ("Confidence", lambda r: r["snapshot"].rotation_confidence), ("Reliable Value", lambda r: r["snapshot"].reliable_value)], rows)
    add_table("Movers", [("Player", lambda r: r["player"].full_name), ("7D Value Change", lambda r: r["history"]["7D"]["delta_value"]), ("7D Price Change", lambda r: r["history"]["7D"]["delta_price"]), ("7D Rank Change", lambda r: r["history"]["7D"]["delta_rank"])], rows)

    fixtures = db.scalars(select(Fixture).order_by(Fixture.event, Fixture.kickoff_time)).all()
    fixture_sheet = workbook.create_sheet("Fixtures")
    fixture_sheet.append(["Gameweek", "Home Team", "Away Team", "Kickoff", "Finished"])
    for fixture in fixtures:
        fixture_sheet.append([_excel_value(value) for value in [fixture.event, fixture.team_h, fixture.team_a, fixture.kickoff_time, fixture.finished]])
    changes = db.scalars(select(SchemaChange).order_by(SchemaChange.detected_at.desc())).all()
    schema_sheet = workbook.create_sheet("Schema Changes")
    schema_sheet.append(["Detected", "Category", "Change", "Field"])
    for change in changes:
        schema_sheet.append([_excel_value(value) for value in [change.detected_at, change.category, change.change_type, change.field_name]])
    guide = workbook.create_sheet("Guide")
    guide.append(["Metric", "Definition"])
    guide.append(["Value", "Total FPL points divided by current price in millions."])
    guide.append(["Perfect Pick", "Points per team match, nudged 15% towards the last three matches, scaled by recent minutes and flagged availability. In points per match. See docs/METRICS.md."])
    guide.append(["Reliable Value", "Value reduced by minutes, start share, and sample-size reliability."])
    guide.append(["Past Seasons sheet", "Completed seasons only. Avg P/90, Spread and Consistency use seasons of 900+ minutes; the pooled measures count from a player's first such season. Starts before 2022/23 are inferred from minutes."])
    guide.append(["Forward Value", "Heuristic projected points over the configured fixture window divided by price."])
    guide.append(["Rotation Risk", "Transparent 0-100 estimate; higher means less secure starts/minutes."])
    guide.append(["Expected Minutes", "Observed minutes per team match, weighted by start share and availability."])
    guide.append(["Empty cell", "The metric has no value: its inputs are not available yet. An empty cell is not a zero."])
    guide.append(["0", "A measured zero. The calculation ran and the answer was zero."])

    sheet.freeze_panes = "A2"

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
